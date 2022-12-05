from openpyxl.chart.bar_chart import BarChart 
from openpyxl.chart.reference import Reference 
from openpyxl.chart.line_chart import LineChart 
 

from openpyxl import load_workbook 

wb = load_workbook("sample.xlsx")
ws = wb.active


# B2:C11 까지의 데이터를 차트로 생성
# bar_value = Reference(ws, min_row=2, max_row=11, min_col=2, max_col=3 )

# bar_chart = BarChart()

# bar_chart.add_data(bar_value)

# ws.add_chart(bar_chart,"E1")



#  B1:C11 까지의 데이터 차트 만들기
Line_value = Reference(ws, min_row=1, max_row=11, min_col=2, max_col=3 )
#차트 생성
Line_chart = LineChart()
# 계열 > 영어 , 수학
Line_chart.add_data(Line_value, titles_from_data=True) 

# 차트 타이틀 입력
Line_chart.title = "성적표"
# 미리 정의된 스타일을 적용, 상ㅇ자가 개별 지정도 가능
Line_chart.style = 20
# Y축의 제목
Line_chart.y_axis.title = "점수"
# X축의 제목
Line_chart.x_axis.title = "번호"


ws.add_chart(Line_chart,"E1")  # type: ignore


wb.save("sample_chart.xlsx")

